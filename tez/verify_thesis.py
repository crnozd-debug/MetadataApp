# -*- coding: utf-8 -*-
"""tezbeta.docx dosyasinin icerik dogrulamasi"""
from docx import Document
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = Document(r'c:\metadataapp\tez\tezbeta.docx')

print(f"Toplam paragraf: {len(doc.paragraphs)}")
print(f"Toplam section: {len(doc.sections)}")
print()

# Sayfa ayarlarini kontrol et
for i, sec in enumerate(doc.sections):
    print(f"Section {i}: {sec.page_width.cm:.1f}x{sec.page_height.cm:.1f} cm, "
          f"Margins: top={sec.top_margin.cm:.1f} bot={sec.bottom_margin.cm:.1f} "
          f"left={sec.left_margin.cm:.1f} right={sec.right_margin.cm:.1f}")

print("\n=== TUM PARAGRAFLAR (ilk 80 karakter) ===")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text:
        style = p.style.name if p.style else 'None'
        bold = 'B' if (p.runs and p.runs[0].bold) else ' '
        size = ''
        if p.runs and p.runs[0].font.size:
            size = f'{p.runs[0].font.size.pt:.0f}pt'
        align = str(p.alignment) if p.alignment else 'None'
        print(f"  [{i:3d}] [{style:15s}] [{bold}] [{size:5s}] {text[:80]}")
